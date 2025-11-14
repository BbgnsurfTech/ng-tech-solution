#!/usr/bin/env python3
"""
AgroConnect360 - Farmer & User Stakeholder Presentation
Creates a user-friendly presentation for farmers and end-users
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_farmer_presentation():
    """Create a farmer/user stakeholder presentation"""

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Define color scheme
    PRIMARY_GREEN = RGBColor(46, 125, 50)
    SECONDARY_ORANGE = RGBColor(245, 124, 0)
    ACCENT_TEAL = RGBColor(0, 137, 123)
    WHITE = RGBColor(255, 255, 255)
    DARK_GRAY = RGBColor(33, 33, 33)

    def add_title_slide(title, subtitle=""):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = PRIMARY_GREEN

        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1))
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(54)
        title_para.font.bold = True
        title_para.font.color.rgb = WHITE
        title_para.alignment = PP_ALIGN.CENTER

        if subtitle:
            subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.7), Inches(9), Inches(0.8))
            subtitle_frame = subtitle_box.text_frame
            subtitle_frame.text = subtitle
            subtitle_para = subtitle_frame.paragraphs[0]
            subtitle_para.font.size = Pt(28)
            subtitle_para.font.color.rgb = WHITE
            subtitle_para.alignment = PP_ALIGN.CENTER

        return slide

    def add_content_slide(title, content_items, use_bullets=True):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Title bar
        title_shape = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(10), Inches(1))
        title_shape.fill.solid()
        title_shape.fill.fore_color.rgb = PRIMARY_GREEN
        title_frame = title_shape.text_frame
        title_frame.text = title
        title_frame.paragraphs[0].font.size = Pt(36)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = WHITE
        title_frame.margin_left = Inches(0.5)
        title_frame.margin_top = Inches(0.2)

        # Content
        content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(5.7))
        text_frame = content_box.text_frame
        text_frame.word_wrap = True

        for i, item in enumerate(content_items):
            if i > 0:
                text_frame.add_paragraph()
            p = text_frame.paragraphs[i]
            p.text = item
            p.font.size = Pt(22)
            p.font.color.rgb = DARK_GRAY
            p.space_after = Pt(14)
            if use_bullets:
                p.level = 0

        return slide

    # Slide 1: Title
    add_title_slide(
        "AgroConnect360",
        "Your Partner in Successful Farming 🌾"
    )

    # Slide 2: Welcome
    add_content_slide(
        "Welcome to AgroConnect360! 👋",
        [
            "We understand farming is hard work, and we're here to help!",
            "",
            "AgroConnect360 is your all-in-one farming assistant that:",
            "✅ Helps you grow better crops",
            "✅ Connects you directly to buyers",
            "✅ Gives you access to loans and financial services",
            "✅ Provides weather updates and farming advice",
            "",
            "Everything you need, right on your phone! 📱"
        ],
        use_bullets=False
    )

    # Slide 3: Your Challenges, Our Solutions
    add_content_slide(
        "We Understand Your Challenges",
        [
            "❌ CHALLENGE: Hard to find buyers who pay fair prices",
            "✅ SOLUTION: Direct marketplace to sell your crops at better prices",
            "",
            "❌ CHALLENGE: Difficult to get loans from banks",
            "✅ SOLUTION: Quick loans on your phone, approved in minutes",
            "",
            "❌ CHALLENGE: Don't know when to plant or harvest",
            "✅ SOLUTION: Weather forecasts and expert farming advice",
            "",
            "❌ CHALLENGE: Input dealers too far or too expensive",
            "✅ SOLUTION: Order seeds, fertilizers online with delivery"
        ],
        use_bullets=False
    )

    # Slide 4: How AgroConnect360 Helps You
    add_content_slide(
        "How We Help You Succeed 🚀",
        [
            "🌾 BETTER FARMING: Track your crops, get planting reminders, expert advice",
            "",
            "💰 MORE MONEY: Sell directly to buyers, no middlemen taking your profits",
            "",
            "📱 EASY ACCESS: Everything on your phone - no need to travel far",
            "",
            "🌤️ WEATHER ALERTS: Know when rain is coming, plan your farming activities",
            "",
            "💳 QUICK LOANS: Get money for farming when you need it",
            "",
            "📚 LEARN & GROW: Training videos, tips from expert farmers"
        ],
        use_bullets=False
    )

    # Slide 5: Main Features
    add_content_slide(
        "What You Can Do with AgroConnect360",
        [
            "📊 FARM MANAGEMENT",
            "  • Record all your crops and farms in one place",
            "  • Track planting dates, growth stages, and expected harvest",
            "  • Get reminders for watering, fertilizing, and harvesting",
            "",
            "🛒 MARKETPLACE",
            "  • Sell your produce directly to buyers",
            "  • Buy quality seeds, fertilizers, and equipment",
            "  • See fair market prices before selling",
            "",
            "💰 FINANCIAL SERVICES",
            "  • Save money in your digital wallet",
            "  • Get loans from ₦10,000 to ₦5,000,000",
            "  • Pay easily using your phone"
        ],
        use_bullets=False
    )

    # Slide 6: Success Stories
    add_content_slide(
        "Success Stories from Farmers Like You 🌟",
        [
            "👨‍🌾 Adeola, Cocoa Farmer (Oyo State):",
            "\"I used to sell my cocoa for ₦800/kg to middlemen. With AgroConnect360, I found buyers paying ₦1,200/kg. I made ₦400,000 more this season!\"",
            "",
            "👩‍🌾 Chioma, Cassava Farmer (Ogun State):",
            "\"The weather alerts saved my crops! I harvested before heavy rains that damaged my neighbor's farm. Plus, I got a ₦200,000 loan in just 2 days!\"",
            "",
            "👨‍🌾 Musa, Rice Farmer (Niger State):",
            "\"The farming tips helped me increase my yield by 40%. The app also connected me to an input dealer with cheaper fertilizers.\""
        ],
        use_bullets=False
    )

    # Slide 7: Getting Started
    add_content_slide(
        "How to Get Started - It's Easy! ✨",
        [
            "1️⃣ DOWNLOAD THE APP",
            "   • Available on Google Play Store and App Store",
            "   • Works on any smartphone",
            "",
            "2️⃣ CREATE YOUR ACCOUNT",
            "   • Use your phone number",
            "   • Takes less than 2 minutes",
            "",
            "3️⃣ SET UP YOUR FARM",
            "   • Add your farm location and size",
            "   • Record your crops",
            "",
            "4️⃣ START ENJOYING BENEFITS",
            "   • Browse marketplace, apply for loans, get weather updates"
        ],
        use_bullets=False
    )

    # Slide 8: Pricing
    add_content_slide(
        "How Much Does It Cost? 💵",
        [
            "📱 FREE TO USE!",
            "  • Download and use the app for FREE",
            "  • No monthly subscription fees",
            "",
            "💰 WE ONLY CHARGE WHEN YOU MAKE MONEY:",
            "  • Small 3% fee when you sell crops (only when you get paid)",
            "  • 15% interest on loans (lowest in the market)",
            "  • Free wallet and money transfers",
            "",
            "🎁 SPECIAL OFFER:",
            "  • First 3 months - NO FEES on marketplace sales!",
            "  • First loan - ONLY 10% interest rate!"
        ],
        use_bullets=False
    )

    # Slide 9: Support & Training
    add_content_slide(
        "We're Here to Help You! 🤝",
        [
            "📞 CUSTOMER SUPPORT:",
            "  • Call or WhatsApp: Available 7 days a week",
            "  • In-app chat support",
            "  • Help in English, Yoruba, Hausa, and Igbo",
            "",
            "🎓 FREE TRAINING:",
            "  • Video tutorials on how to use the app",
            "  • Farming best practices",
            "  • Weekly tips and advice",
            "",
            "👥 FARMER COMMUNITY:",
            "  • Connect with other farmers",
            "  • Share experiences and learn from each other",
            "  • Join farmer cooperatives"
        ],
        use_bullets=False
    )

    # Slide 10: Safety & Security
    add_content_slide(
        "Your Money & Information Are Safe 🔒",
        [
            "✅ SECURE PAYMENTS:",
            "  • Bank-level security for all transactions",
            "  • Your money is protected",
            "",
            "✅ PRIVACY PROTECTED:",
            "  • Your personal information stays private",
            "  • We never sell your data",
            "",
            "✅ VERIFIED BUYERS & SELLERS:",
            "  • All users are verified",
            "  • Rate and review other users",
            "",
            "✅ LICENSED & REGULATED:",
            "  • Approved by Nigerian regulatory authorities",
            "  • Following all banking and data protection rules"
        ],
        use_bullets=False
    )

    # Slide 11: What Farmers Are Saying
    add_content_slide(
        "What Other Farmers Say About Us ⭐⭐⭐⭐⭐",
        [
            "\"Best farming app I've ever used. Easy to understand and very helpful!\"",
            "- Tunde, Maize Farmer",
            "",
            "\"I got my loan approved in 30 minutes. Used it to buy fertilizer and my crops are doing great!\"",
            "- Blessing, Vegetable Farmer",
            "",
            "\"The weather alerts are always accurate. Saved my harvest twice!\"",
            "- Ibrahim, Rice Farmer",
            "",
            "\"Finally, an app that speaks my language and understands our farming needs.\"",
            "- Ngozi, Cassava Farmer"
        ],
        use_bullets=False
    )

    # Slide 12: Join Us Today
    add_content_slide(
        "Join 10,000+ Successful Farmers Today! 🎉",
        [
            "✅ Increase your farm income by up to 40%",
            "✅ Save time and money",
            "✅ Get expert advice whenever you need it",
            "✅ Access to loans and financial services",
            "✅ Connect with buyers across Nigeria",
            "",
            "📲 DOWNLOAD NOW:",
            "  • Google Play Store: Search 'AgroConnect360'",
            "  • Apple App Store: Search 'AgroConnect360'",
            "",
            "🎁 Sign up this month and get 3 months FREE marketplace fees!"
        ],
        use_bullets=False
    )

    # Slide 13: Contact & Closing
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = PRIMARY_GREEN

    text_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2.5))
    text_frame = text_box.text_frame
    text_frame.text = "Farming Made Simple,\nSuccess Made Possible! 🌾"
    p = text_frame.paragraphs[0]
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    contact_box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(8), Inches(1.8))
    contact_frame = contact_box.text_frame
    contact_frame.text = "📧 support@agroconnect360.ng\n📱 +234-XXX-XXXX-XXX (Call/WhatsApp)\n🌐 www.agroconnect360.ng"
    for p in contact_frame.paragraphs:
        p.font.size = Pt(24)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    # Save presentation
    filename = "presentations/AgroConnect360_Farmer_Presentation.pptx"
    prs.save(filename)
    print(f"✅ Farmer presentation created: {filename}")
    return filename

if __name__ == "__main__":
    create_farmer_presentation()
